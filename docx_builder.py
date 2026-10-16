# -*- coding: utf-8 -*-
"""Word 生成层：md_parser 的块树 + formula_converter 的 OMML → .docx。

这一层的边界：**只认块树，不认 Markdown，也不认 MathML**。它做五件事：

1. 建样式（正文宋体 12pt、标题黑体并挂到 Heading 1/2/3、代码用等宽字体）。
2. 逐块落笔：段落 / 标题 / 块级公式 / 表格 / 列表 / 引用 / 代码块 / 分隔线。
3. 把 OMML 元素挂进 **段落和表格单元格**——公式在 Word 里是原生对象，不是图片。
4. 汇总一份回执：转成几条、失败几条（原文 + 中文原因）、图片嵌进几处 / 没成几处（每处一句原因）、
   超链接几个。
5. 按 :class:`MergeLayout` 摆合并稿的前置页：封面、目录（Word 的 TOC 域）、文件之间换页。
   **哪几个块前面要换页、封面上写什么字，是策略层（gui/merger.py）定好的**——
   这一层只认块树与下标，不认 Markdown，也不去数文件。

四个"不这样做 Word 就会出错或者看着不对"的坑，都写在这当档：

* **中文字体要写 ``w:eastAsia``**。``run.font.name`` / ``style.font.name`` 只写
  ``w:ascii`` 与 ``w:hAnsi``，那两格管西文；中文取的是 ``w:eastAsia``，不写就用主题字体，
  宋体黑体压根不生效。字号同理：``font.size`` 只写 ``w:sz``，``w:szCs`` 要自己跟上
  （python-docx 的 ``CT_RPr`` 压根没声明 ``szCs``）。
* **主题属性要删**。模板里 Heading 样式的 ``w:rFonts`` 四格全是 ``*Theme``
  （``asciiTheme="majorHAnsi"`` 等），``w:color`` 带 ``themeColor="accent1"``；
  规范规定主题属性优先于显式 val，只 set 不删就等于设了一格不被读的，标题照旧是蓝色 Calibri Light。
* **子元素次序是 schema 的一部分**。OOXML 里 ``w:pPr``/``w:rPr``/``w:tblPr`` 都是**序列类型**，
  次序摆错 Word 直接拒开这个包。python-docx 自己那批 setter 排得对，本层手写的那几格
  （``w:tblBorders``、``w:shd``、``w:pBdr``、``w:szCs``）必须走 :func:`_insert_in_order`
  ——实测 ``add_table()`` 生成的 ``w:tblPr`` 里已经有 ``w:tblLook``，它排在 ``w:tblBorders`` **之后**，
  所以"append 到末尾"这种写法会稳定产出一个打不开的 docx。
* **``w:hyperlink`` 与 OMML 是自己接上去的**（python-docx 没有公开 API）。
  所以本层每格文字都走 :meth:`DocxBuilder._append_run`（把 ``w:r`` 直接 append 到 ``w:p`` 末尾），
  这样"文字—公式—链接—公式"的次序就是块树里的次序，不受各 setter 的插入规则影响。

失败策略（业务方 2026-10-14 定）：一条公式转不出 OMML 时**不中断整份文档**，
把 LaTeX 原文带定界符写回去并上黄色底纹，同时记进回执；``--strict`` 才当场抛错。
理由是真实用法是"手里一份几百条公式的稿子"——为一条写坏的公式出一份残缺文档，
比出一份带标记的文档更难修。
"""

from __future__ import annotations

import xml.etree.ElementTree as _StdET
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Optional, Sequence, Union

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.image.exceptions import (
    InvalidImageStreamError,
    UnexpectedEndOfFileError,
    UnrecognizedImageError,
)
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, RGBColor
from docx.shape import InlineShape
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from lxml import etree

import md_parser as mp
from formula_converter import (
    MATHML_NS,
    MATH_NS,
    WORD_NS,
    ConversionResult,
    FormulaConversionError,
    FormulaConverter,
)

# --------------------------------------------------------------------------
# 排版常量（能被命令行覆盖的是那四个字体/字号，这里收的是结构性的东西）
# --------------------------------------------------------------------------

BODY_FONT_DEFAULT = "宋体"
BODY_SIZE_DEFAULT = 12.0
HEADING_FONT_DEFAULT = "黑体"
CODE_FONT_DEFAULT = "Consolas"

#: Markdown 第 N 级标题 → Word 内建样式名。4~6 级规格里没点名，按同一思路顺延。
HEADING_STYLES: dict[int, str] = {
    1: "Heading 1",
    2: "Heading 2",
    3: "Heading 3",
    4: "Heading 4",
    5: "Heading 5",
    6: "Heading 6",
}

#: 各级标题相对正文字号的**增量**（pt）。写死绝对号的话 --body-size 一改就全乱。
HEADING_STEPS: dict[int, float] = {1: 5.0, 2: 3.0, 3: 1.5, 4: 0.5, 5: 0.0, 6: 0.0}

#: Markdown 表格的对齐声明 → Word 段落对齐。``none`` 不设，交给 Word 默认。
ALIGNMENTS: dict[str, Optional[WD_ALIGN_PARAGRAPH]] = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "none": None,
}

BULLET_GLYPH = "•"
#: 一层缩进（磅）。列表、引用、代码块共用一个数，改一处就够了。
INDENT_PT = 21.0
#: 超链接颜色：Word 自己的链接蓝。画成黑色下划线会被读成普通文字。
LINK_COLOR = RGBColor(0x05, 0x63, 0xC1)
#: 转换失败那条 LaTeX 的高亮色——"这里要人来看"的唯一视觉信号。
FALLBACK_HIGHLIGHT = WD_COLOR_INDEX.YELLOW
CODE_FILL = "F2F2F2"
HEADER_FILL = "D9D9D9"
QUOTE_COLOR = RGBColor(0x59, 0x59, 0x59)
QUOTE_BORDER_COLOR = "A6A6A6"
BORDER_COLOR = "auto"
BORDER_SIZE_EIGHTHS = 4  # w:sz 的单位是 1/8 磅，4 = 0.5 磅

#: 图片**嵌不进 Word 时**留在正文上的那句占位。口径（业务方 2026-10-16 改）：
#: 盘上读得到的图就**真的嵌进去**，读不到的留一句带原因的话——
#: 悄悄丢掉图片是丢信息，而占位里写着文件名和为什么没成，人一眼就能找到原图去补。
IMAGE_PLACEHOLDER = "【图片没嵌进去：{name}——{reason}】"
#: 四类"没成"的原因各一句。为什么分开写：这四件事的下一步完全不同
#: （换成本地文件 / 把地址补上 / 改图格式 / 去那个位置看一眼）。
IMAGE_SKIP_REMOTE = "这不是本机的文件（远程或内联的地址，本工具不去取）"
IMAGE_SKIP_NO_ADDRESS = "这一条没写图片地址"
IMAGE_SKIP_MISSING = "这台机器上没找到这个文件（查的位置：%s）"
IMAGE_SKIP_FORMAT = "这个文件读不动——Word 里能放的是 PNG/JPEG/GIF/TIFF/BMP 与 EMF/WMF"
IMAGE_SKIP_FAILED = "这个文件读不动（%s）"

# --- 合并稿的那一叠（封面 / 目录域 / 文件之间换页）-----------------------------

#: 目录标题那四个字。**不挂 Heading 样式**——挂了目录就会把"目录"自己列进目录里。
TOC_HEADING_DEFAULT = "目录"
#: 目录域收录到哪一级（`\o` 那一档）。合并稿传的不是这个默认值：
#: 每个文件的标题都降了一级，原来的 `#` 落在第 2 级，收到 3 就把它底下的 `###` 漏了。
TOC_LEVELS_DEFAULT = "1-3"
#: 域代码。`\h` 让条目能点、`\z` 网页视图不显示页码、`\u` 按大纲级取标题
#: （Word 内建的 Heading 样式自带 `w:outlineLvl`，所以本层挂上样式就等于进了目录）。
TOC_INSTRUCTION = ' TOC \\o "{levels}" \\h \\z \\u '
#: 域的**缓存结果**。Word 一刷新就把这一行换掉，所以这句"要去更新域"只在该刷新没生效时还留在页上。
TOC_PLACEHOLDER = ("（这一栏是 Word 的目录域。打开这个文件时会自动刷新；"
                   "要是还看得见这一行，全选后按 F9，或右键「更新域」。）")
#: 页码是 Word 排版之后现算的，不在本层的能力范围内——这一句写进回执，别说"目录已生成"。
TOC_NOTE = ("目录那一栏是 Word 的域（TOC \\o…）：页码要 Word 排版之后才数得出来，"
            "所以文件里先放占位并让它打开时自动刷新；没刷新就在 Word 里「更新域」")
#: 封面标题离页顶的距离（磅）。用 ``space_before`` 而不是连摞几个空段——
#: 摞空段那种写法换一台打印机/换一个行距设置，标题就高一行低一行。
COVER_TITLE_SPACE_PT = 96.0
COVER_TITLE_EXTRA_PT = 12.0
COVER_LINE_EXTRA_PT = 3.0
COVER_LINE_SPACE_PT = 30.0


@dataclass(frozen=True)
class CoverPage:
    """封面那三行。**文本在这里定稿**，本层不做日期格式化也不数文件——
    那些是合并策略（gui/merger.py）的事，这一层只管"摆进 Word"。"""

    title: str = ""
    date_text: str = ""
    file_count_text: str = ""


@dataclass(frozen=True)
class MergeLayout:
    """一份合并稿的排版声明：封面、目录、以及**哪些块前面要换页**。

    ``section_starts`` 装的是块树的**下标**（不是文件名、不是行号）：策略层先把各文件的块
    拼成一棵，再告诉这里"第 3 个块前面换页"。为什么不带文件名——本层不认 Markdown，
    给它文件名等于让它去猜哪个块属于哪个文件（§文档开头那条边界）。
    ``toc_heading`` 空串 = 这一份不要目录，不是"目录标题为空"。
    """

    cover: Optional[CoverPage] = None
    toc_heading: str = ""
    toc_levels: str = TOC_LEVELS_DEFAULT
    section_starts: tuple[int, ...] = ()

# --- 各父元素的子元素次序（只列本层会碰到的那一族；不在名单里的标签按"排在它之后"处理）
PPR_SEQ: tuple[str, ...] = (
    "w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr",
    "w:widowControl", "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd",
    "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
    "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
    "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
    "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
    "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr",
    "w:sectPr", "w:pPrChange",
)
RPR_SEQ: tuple[str, ...] = (
    "w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps", "w:smallCaps",
    "w:strike", "w:dstrike", "w:outline", "w:shadow", "w:emboss", "w:noProof",
    "w:snapToGrid", "w:vanish", "w:webHidden", "w:color", "w:spacing", "w:w", "w:kern",
    "w:position", "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd",
    "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout",
    "w:specVanish", "w:oMath",
)
TCPR_SEQ: tuple[str, ...] = (
    "w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders", "w:shd",
    "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark",
)
TRPR_SEQ: tuple[str, ...] = (
    "w:cnfStyle", "w:divId", "w:gridBefore", "w:gridAfter", "w:wBefore", "w:wAfter",
    "w:cantSplit", "w:trHeight", "w:tblHeader", "w:tblCellSpacing", "w:jc", "w:hidden",
)
TBLPR_SEQ: tuple[str, ...] = (
    "w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize",
    "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd",
    "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
)
#: ``word/settings.xml`` 的根也是序列类型。名单只列本层会碰到的那一族：
#: ``w:updateFields`` 必须排在 ``w:compat`` **之前**（实测默认模板的子元素次序是
#: zoom → proofState → defaultTabStop → characterSpacingControl → savePreviewPicture →
#: compat → rsids），次序错了不是"目录不出"而是整个包被判无效。
SETTINGS_SEQ: tuple[str, ...] = (
    "w:writeProtection", "w:view", "w:zoom", "w:removePersonalInformation",
    "w:proofState", "w:defaultTabStop", "w:characterSpacingControl",
    "w:savePreviewPicture", "w:updateFields", "w:hdrShapeDefaults", "w:footnotePr",
    "w:endnotePr", "w:compat", "w:docVars", "w:rsids", "w:mathPr", "w:themeFontLang",
    "w:clrSchemeMapping",
)


# --------------------------------------------------------------------------
# 回执
# --------------------------------------------------------------------------


@dataclass
class BuilderOptions:
    """可覆盖的排版设置（对应 --body-font / --body-size / --heading-font）。"""

    body_font: str = BODY_FONT_DEFAULT
    body_size: float = BODY_SIZE_DEFAULT
    heading_font: str = HEADING_FONT_DEFAULT
    code_font: str = CODE_FONT_DEFAULT

    def heading_size(self, level: int) -> float:
        return self.body_size + HEADING_STEPS.get(level, 0.0)


@dataclass
class BuildReport:
    """一次生成的回执。``document`` 是成品，其余是给命令行报告用的数。"""

    document: Document
    options: BuilderOptions = field(default_factory=BuilderOptions)
    math_ok: int = 0
    math_failed: int = 0
    links: int = 0
    images: int = 0
    #: 真嵌进 Word 的处数与没成的处数。**两个都单独给**：只念"图片 3 处"会把"有一张没找到"
    #: 演成没事（同"未组价的行也要单独报出来"那一条铁律）。
    images_embedded: int = 0
    images_skipped: int = 0
    tables: int = 0
    blocks: int = 0
    failures: list[ConversionResult] = field(default_factory=list)
    parse_warnings: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    @property
    def math_total(self) -> int:
        return self.math_ok + self.math_failed

    def add_failure(self, result: ConversionResult) -> None:
        self.math_failed += 1
        self.failures.append(result)

    def lines(self) -> list[str]:
        """给人读的报告正文（文件名与路径由 main 拼）。"""
        out = [
            "块 %d 个｜表格 %d 张｜公式 %d 条（转成 %d，失败 %d）｜超链接 %d 个｜图片 %d 处（嵌进 %d，没成 %d）"
            % (
                self.blocks,
                self.tables,
                self.math_total,
                self.math_ok,
                self.math_failed,
                self.links,
                self.images,
                self.images_embedded,
                self.images_skipped,
            )
        ]
        out.extend("  · " + text for text in self.parse_warnings)
        out.extend("  · " + text for text in self.notes)
        if self.failures:
            out.append("转换失败的公式（正文里已用原文 + 黄色底纹标出，可直接在 Word 里改）：")
            out.extend("  × " + item.failure_report() for item in self.failures)
        return out

    def save(self, path: str | Path) -> Path:
        """落盘。父目录不存在就先建（命令行里指到没建过的子目录是很常见的写法）。"""
        target = Path(path)
        if target.parent and not target.parent.exists():
            target.parent.mkdir(parents=True, exist_ok=True)
        self.document.save(str(target))
        return target


# --------------------------------------------------------------------------
# 落笔处的抽象：正文 or 表格单元格
# --------------------------------------------------------------------------

Host = Union[Document, _Cell]


class _Host:
    """"能往这里放段落和表"的地方。

    正文与单元格只差在第一段：``_Cell`` 自带一个空段落（不用的话单元格里就多一个空行），
    ``Document`` 的 body 一段都没有。另外这里会复用"还没写过东西的那一段"，
    免得块树里相邻两个节点各开一段、文档中间凭空多出一堆空行。
    """

    __slots__ = ("target", "_current", "_used")

    def __init__(self, target: Host) -> None:
        self.target = target
        self._current: Optional[Paragraph] = None
        self._used = False

    @property
    def paragraphs(self) -> Sequence[Paragraph]:
        """这一处已有的段落（正文取 body 直属段落，单元格取它自己的）。"""
        return self.target.paragraphs

    def paragraph(self) -> Paragraph:
        current = self._current
        if current is not None and not _has_content(current):
            return current
        if isinstance(self.target, _Cell):
            if not self._used:
                self._used = True
                current = self.target.paragraphs[0]
            else:
                current = self.target.add_paragraph()
        else:
            current = self.target.add_paragraph()
        self._current = current
        return current

    def table(self, rows: int, cols: int) -> Table:
        if isinstance(self.target, _Cell):
            return self.target.add_table(rows, cols)
        return self.target.add_table(rows, cols)


def _has_content(paragraph: Paragraph) -> bool:
    """段落里除了 ``w:pPr`` 还有没有别的东西——有就不能当"空段"再复用。"""
    return any(child.tag != qn("w:pPr") for child in paragraph._p)


# --------------------------------------------------------------------------
# OMML 挂载
# --------------------------------------------------------------------------

#: document.xml 里用 m:/w:/mml: 而不是 lxml 自动生成的 ns0:/ns1:
_PREFIXES: tuple[tuple[str, str], ...] = (
    ("m", MATH_NS),
    ("w", WORD_NS),
    ("mml", MATHML_NS),
)


def _omml_for_word(element: etree._Element) -> etree._Element:
    """把 formula_converter 产出的元素换成带 ``m:``/``w:`` 前缀的版本，再交给 python-docx 的解析器。

    Word 只认命名空间 URI、不认前缀，所以这一步**不影响能不能打开**，影响的是可读性：
    对账时要人解压 docx 看 ``word/document.xml``（这正是"公式到底进没进文件"的现场），
    ``<ns0:oMathPara>`` 与 ``<m:oMathPara>`` 差一口气。
    """
    tree = _StdET.fromstring(etree.tostring(element))
    for prefix, uri in _PREFIXES:
        _StdET.register_namespace(prefix, uri)
    return parse_xml(_StdET.tostring(tree))


def count_omml(path: str | Path) -> tuple[int, int]:
    """读回落盘的 .docx，数 ``word/document.xml`` 里的 ``m:oMath`` 与 ``m:oMathPara``。

    这是"文件里真的有原生公式对象"的唯一证据：内存里的树说得再好听，
    落盘那一步被丢掉了才是事故。按命名空间 URI 数，不按标签字符串数（前缀写成什么不该影响计数）。
    """
    with zipfile.ZipFile(str(path)) as package:
        xml = package.read("word/document.xml")
    root = etree.fromstring(xml)
    inline = len(root.findall(".//{%s}oMath" % MATH_NS))
    display = len(root.findall(".//{%s}oMathPara" % MATH_NS))
    return inline, display


# ==========================================================================
# 生成器
# ==========================================================================


class DocxBuilder:
    """块树 → Word 文档。

    ``converter`` 由调用方构造（它只管公式，不知道 Word 的存在）；
    ``strict=True`` 时第一条转不出 OMML 的公式就抛 :class:`FormulaConversionError`。
    """

    def __init__(
        self,
        converter: FormulaConverter,
        options: Optional[BuilderOptions] = None,
        strict: bool = False,
    ) -> None:
        self.converter = converter
        self.options = options or BuilderOptions()
        self.strict = strict

    # ---------------------------------------------------------------- 对外入口
    def build(self, parsed: mp.ParseResult, layout: Optional[MergeLayout] = None) -> BuildReport:
        """整棵块树落进一份新文档，返回回执（``report.document`` 就是成品）。

        ``layout=None`` 是"这一份就一个文件"的那一路（命令行、单独输出）：
        封面、目录、分页一个都不写，**与加这一格之前的产物逐字节相同**。
        """
        document = Document()
        report = BuildReport(document=document, options=self.options)
        report.parse_warnings = list(parsed.warnings)
        self._setup_styles(document, report)

        host = _Host(document)
        if layout is not None:
            self._write_front_matter(host, layout, report)
        starts = frozenset(layout.section_starts) if layout is not None else frozenset()
        for index, block in enumerate(parsed.blocks):
            if index in starts:
                self._write_page_break(host)
            self._write_block(host, block, report)
            report.blocks += 1
        if not parsed.blocks:
            host.paragraph()
            report.notes.append("输入文件解析出来是空的，已写一份只有空段落的文档")

        # 正文末尾不该压在表格上（Word 会自己补一段，但补在哪儿不如我们定）。
        body = [child for child in document.element.body if child.tag != qn("w:sectPr")]
        if body and body[-1].tag == qn("w:tbl"):
            document.add_paragraph()
        return report

    def build_file(self, parsed: mp.ParseResult, path: str | Path) -> BuildReport:
        """一步到位：解析结果 → 落盘。命令行走这一条。"""
        report = self.build(parsed)
        report.save(path)
        return report

    # ------------------------------------------------------------------ 样式
    def _setup_styles(self, document: Document, report: BuildReport) -> None:
        """字体字号写进**样式**而不是逐段设 run：一份稿子上千段，逐段设既慢又让
        document.xml 胖一倍，而人在 Word 里改样式本来就该全文跟着变。"""
        normal = self._style(document, "Normal")
        if normal is not None:
            _apply_font(normal, self.options.body_font)
            _apply_size(normal, self.options.body_size)

        for level, style_name in HEADING_STYLES.items():
            style = self._style(document, style_name)
            if style is None:
                report.notes.append("模板里没有样式「%s」，第 %d 级标题改用正文样式 + 加粗" % (style_name, level))
                continue
            _apply_font(style, self.options.heading_font)
            _apply_size(style, self.options.heading_size(level))
            style.font.bold = True
            _make_color_black(style)

    @staticmethod
    def _style(document: Document, name: str):
        """按名字取样式，取不到交回 None（换过模板时不该整份文档崩掉）。"""
        try:
            return document.styles[name]
        except KeyError:
            return None

    # ---------------------------------------------------------------- 合并稿的前置页
    def _write_front_matter(self, host: _Host, layout: MergeLayout, report: BuildReport) -> None:
        """封面 → 换页 → 目录 → 换页。正文因此一定从新的一页开头。

        这两处分页**不受「文件之间插入分页符」那个开关管**，也刻意不并进去：
        关了它之后正文会紧接在目录下面，那一张目录就没法读——
        开关说的是"各文件之间要不要各起一页"，前置页要不要独占一页本来就不是同一件事。
        """
        if layout.cover is not None:
            self._write_cover(host, layout.cover, report)
            self._write_page_break(host)
        if layout.toc_heading:
            self._write_toc(host, layout, report)
            self._write_page_break(host)

    def _write_cover(self, host: _Host, cover: CoverPage, report: BuildReport) -> None:
        """封面三行：标题、日期、文件数。**都不挂 Heading 样式**——
        挂了封面那几行就进了目录，而"共 3 个文件"这种句子出现在条目里很难看。

        离页顶的距离走 ``space_before`` 而不是连摞几个空段：摞空段那种写法换一台打印机、
        换一个行距设置，标题就高一行低一行，而空段本身还在文档里（选中时会多几个空行）。
        """
        lines = [(text, is_title) for text, is_title in (
            (cover.title, True), (cover.date_text, False), (cover.file_count_text, False),
        ) if text]
        for position, (text, is_title) in enumerate(lines):
            paragraph = host.paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(
                COVER_TITLE_SPACE_PT if position == 0 else COVER_LINE_SPACE_PT
            )
            run = self._append_run(paragraph)
            if is_title:
                _apply_font(run, self.options.heading_font)
                _apply_size(run, self.options.heading_size(1) + COVER_TITLE_EXTRA_PT)
                _set_weight(run, True, False)
            else:
                _apply_font(run, self.options.body_font)
                _apply_size(run, self.options.body_size + COVER_LINE_EXTRA_PT)
            _set_run_text(run, text)

    def _write_toc(self, host: _Host, layout: MergeLayout, report: BuildReport) -> None:
        """目录 = 一个**复合域**（fldChar begin → instrText → separate → 占位 → end）。

        为什么不用更短的 ``w:fldSimple``：域的结果是多段的（一个条目一段），
        而 fldSimple 是"inline 在一个 w:p 里"的东西，Word 读得懂却排不出目录那个样子；
        复合域正是 Word「引用 → 目录」自己写出来的那一套。

        域代码里只写字段名和开关，**页码一个字都不写**：页码是 Word 排版之后现算的，
        本层写死就是假数。所以配套的是 :meth:`_mark_update_fields`（打开时自动刷新）
        与下面那句回执（没刷新时怎么手动更新）。
        """
        heading = host.paragraph()
        run = self._append_run(heading)
        _apply_font(run, self.options.heading_font)
        _apply_size(run, self.options.heading_size(2))
        _set_weight(run, True, False)
        _set_run_text(run, layout.toc_heading)

        paragraph = host.paragraph()
        self._field_char(self._append_run(paragraph), "begin")
        _set_instr_text(
            self._append_run(paragraph), TOC_INSTRUCTION.format(levels=layout.toc_levels)
        )
        self._field_char(self._append_run(paragraph), "separate")
        # separate 与 end 之间是"域的上一次结果"。留着占位而不是留空：Word 没刷新时
        # 页面上还写着这一行，人看得见这里有一栏待更新；留空就只会以为这份文件没有目录。
        _set_run_text(self._append_run(paragraph), TOC_PLACEHOLDER)
        self._field_char(self._append_run(paragraph), "end")

        self._mark_update_fields(report.document)
        report.notes.append(TOC_NOTE)

    @staticmethod
    def _field_char(run: Run, kind: str) -> None:
        """``w:fldChar``：begin / separate / end 三种状态都挂在这个元素上（``w:fldCharType``）。"""
        element = OxmlElement("w:fldChar")
        element.set(qn("w:fldCharType"), kind)
        run._element.append(element)

    @staticmethod
    def _mark_update_fields(document: Document) -> None:
        """``w:updateFields val="true"`` → Word 一打开就重算所有域，目录自己就长出页码。

        这一格排在 ``w:compat`` 之前（见 :data:`SETTINGS_SEQ`），所以走
        :func:`_insert_in_order` 而不是 append。已有的那一格只改值、不重复插。
        """
        settings = document.settings.element
        element = settings.find(qn("w:updateFields"))
        if element is None:
            element = OxmlElement("w:updateFields")
            _insert_in_order(settings, element, SETTINGS_SEQ)
        element.set(qn("w:val"), "true")

    def _write_page_break(self, host: _Host) -> None:
        """分页符单独成段，不并进上一段末尾。

        并进正文那一段会让最后一段多出一个"看不见却占位置"的 run；
        ``w:br type="page"`` 是需求里"文件之间插入分页符"最直白的落法
        （另一条 ``w:pageBreakBefore`` 挂在段落样式上，语义是"这一段起新页"，
        读回来看不出这里原本换了个文件）。
        """
        run = self._append_run(host.paragraph())
        run.add_break(WD_BREAK.PAGE)


    # ------------------------------------------------------------------ 块级
    def _write_block(self, host: _Host, block: mp.Block, report: BuildReport, depth: int = 0) -> None:
        if isinstance(block, mp.Heading):
            self._write_heading(host, block, report, depth)
        elif isinstance(block, mp.Paragraph):
            self._write_inlines(host, block.inline, report, depth=depth)
        elif isinstance(block, mp.BlockMath):
            self._write_block_math(host, block, report, depth)
        elif isinstance(block, mp.CodeBlock):
            self._write_code_block(host, block, report, depth)
        elif isinstance(block, mp.ListBlock):
            self._write_list(host, block, report, depth)
        elif isinstance(block, mp.QuoteBlock):
            self._write_quote(host, block, report, depth)
        elif isinstance(block, mp.ThematicBreak):
            self._write_thematic_break(host, report, depth)
        elif isinstance(block, mp.Table):
            self._write_table(host, block, report, depth)
        else:  # 解析层加了新块而这里没跟上：宁可留一句提醒，不要静默丢内容
            report.notes.append(
                "第 %s 行有一种本工具还不认的块（%s），已跳过"
                % (getattr(block, "line", "?"), type(block).__name__)
            )

    def _write_heading(self, host: _Host, block: mp.Heading, report: BuildReport, depth: int) -> None:
        level = min(max(block.level, 1), 6)
        paragraph = host.paragraph()
        style = self._style(report.document, HEADING_STYLES[level])
        if style is not None:
            paragraph.style = style
        # 样式没挂上时把加粗写进 run，标题至少视觉上还是个标题
        self._write_inlines(
            host, block.inline, report,
            depth=depth, fallback_bold=style is None, styled=style is not None,
        )

    def _write_block_math(self, host: _Host, block: mp.BlockMath, report: BuildReport, depth: int) -> None:
        """独占一段的显示式公式：段落居中 + 外包 ``m:oMathPara``（Word 认这是显示式，
        ``\\sum`` 的上下限才会跑到正上方/正下方）。"""
        paragraph = host.paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_indent(paragraph, depth)
        self._append_math(paragraph, block.latex, display=True, report=report)

    def _write_code_block(self, host: _Host, block: mp.CodeBlock, report: BuildReport, depth: int) -> None:
        """围栏代码块：整块一段、行间用换行连起来，等宽字体 + 浅灰底。

        里面的 ``$`` 在解析层就已经挡在公式识别之外了，这里不必再防。
        """
        paragraph = host.paragraph()
        _set_indent(paragraph, depth)
        _shade_paragraph(paragraph, CODE_FILL)
        paragraph.paragraph_format.space_after = Pt(6)
        lines = block.text.split("\n")
        if lines and not lines[-1]:
            lines.pop()  # 闭合围栏前那个空行不算正文
        for index, line in enumerate(lines):
            run = self._append_run(paragraph)
            _apply_font(run, self.options.code_font)
            _apply_size(run, max(self.options.body_size - 1.0, 8.0))
            if index:
                run.add_break(WD_BREAK.LINE)
            _set_run_text(run, line)

    def _write_list(self, host: _Host, block: mp.ListBlock, report: BuildReport, depth: int) -> None:
        """列表：记号（圆点 / 序号）**写在正文里**，不挂 Word 的列表样式。

        理由：Word 的自动编号要靠 ``word/numbering.xml`` 里的编号定义，光给段落一个
        「List Number」样式名并不会出数字（python-docx 默认模板里那套定义不完整，
        实测开出来是几行光秃秃的段落）。写死记号的代价是"在 Word 里按 Tab 不会自动续号"，
        换来的是**屏幕上看到的编号一定是这份稿子里的那几个数**。
        """
        for index, item in enumerate(block.items, start=block.start):
            marker = "%d." % index if block.ordered else BULLET_GLYPH
            self._write_list_item(host, item, report, depth, marker)

    def _write_list_item(self, host: _Host, item: mp.ListItem, report: BuildReport, depth: int, marker: str) -> None:
        paragraph = host.paragraph()
        _set_indent(paragraph, depth + 1)
        paragraph.paragraph_format.space_after = Pt(0)
        run = self._append_run(paragraph)
        _apply_font(run, self.options.body_font)
        _apply_size(run, self.options.body_size)
        _set_run_text(run, marker + " ")
        self._write_inlines(host, item.inline, report, depth=depth + 1, into=paragraph)
        for child_index, child in enumerate(item.children, start=1):
            child_marker = "%d." % child_index if child.ordered else BULLET_GLYPH
            self._write_list_item(host, child, report, depth + 1, child_marker)

    def _write_quote(self, host: _Host, block: mp.QuoteBlock, report: BuildReport, depth: int) -> None:
        """引用：整块多缩进一层 + 左边一条竖线 + 灰字。

        不用 Word 的「Intense Quote」样式——那一档在换过模板的机器上不保证存在，
        而缩进和边框是自己写的，走到哪台机器都长一个样。
        """
        before = len(host.paragraphs)
        for inner in block.blocks:
            self._write_block(host, inner, report, depth=depth + 1)
        written = list(host.paragraphs)[before:]
        if not written:  # 引用里啥都没写出来（"> " 空着），补一段免得这一整块凭空消失
            written = [host.paragraph()]
        for paragraph in written:
            _quote_style(paragraph, depth + 1)

    def _write_thematic_break(self, host: _Host, report: BuildReport, depth: int) -> None:
        paragraph = host.paragraph()
        _set_indent(paragraph, depth)
        _bottom_rule(paragraph)

    def _write_table(self, host: _Host, block: mp.Table, report: BuildReport, depth: int) -> None:
        rows = block.rows
        if not rows:
            return
        columns = max([block.column_count] + [len(row) for row in rows])
        if columns <= 0:
            return

        table = host.table(len(rows), columns)
        table.autofit = True
        _set_borders(table)
        report.tables += 1

        has_header = block.header_present
        for row_index, row in enumerate(rows):
            is_header = has_header and row_index == 0
            if is_header:
                _repeat_header_row(table.rows[0])
            for column_index in range(columns):
                inlines = row[column_index] if column_index < len(row) else []
                cell = table.cell(row_index, column_index)
                if is_header:
                    _shade_cell(cell, HEADER_FILL)
                self._write_inlines(
                    _Host(cell),
                    inlines,
                    report,
                    align=_alignment_of(block.aligns, column_index),
                    force_bold=is_header,
                )

        # 表后补一个空段：两张表之间没有段落时 Word 会把它们并成一张
        # （列数还对得上就并得更干脆），并完没人看得出这里原本有两张表。
        host.paragraph()

    # ---------------------------------------------------------------- 行内
    def _write_inlines(
        self,
        host: _Host,
        inlines: Iterable[mp.Inline],
        report: BuildReport,
        *,
        depth: int = 0,
        align: Optional[WD_ALIGN_PARAGRAPH] = None,
        force_bold: bool = False,
        fallback_bold: bool = False,
        styled: bool = False,
        into: Optional[Paragraph] = None,
    ) -> None:
        """行内片段列表 → 一段（显示式公式会另起一段并居中）。

        ``into`` 是列表条目那种"段落里已经写了记号"的情况：段落已经有了，
        后续片段接着往同一段里加，不要再开一段。

        ``styled=True`` 是标题那一路：段落已经挂了 Heading 样式，run 上**不再**写
        正文字体与字号——写了就是直接格式，会把样式里的黑体/大字号顶成正文的宋体/正文号，
        而"标题用黑体"正是需求里点名的那条。
        """
        paragraph = into if into is not None else host.paragraph()
        if align is not None:
            paragraph.alignment = align
        if depth:
            _set_indent(paragraph, depth)

        for node in inlines:
            if isinstance(node, mp.MathRun) and node.display:
                # 解析层交来的信号：这个公式要在**这一格内**独占一段、居中
                if _has_content(paragraph):
                    paragraph = host.paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self._append_math(paragraph, node.latex, display=True, report=report, styled=styled)
                paragraph = host.paragraph()
                if align is not None:
                    paragraph.alignment = align
                continue
            self._write_inline(
                paragraph, node, report,
                force_bold=force_bold, fallback_bold=fallback_bold, styled=styled,
            )

    def _write_inline(
        self,
        paragraph: Paragraph,
        node: mp.Inline,
        report: BuildReport,
        *,
        force_bold: bool = False,
        fallback_bold: bool = False,
        styled: bool = False,
    ) -> None:
        bold = force_bold or fallback_bold
        if isinstance(node, mp.TextRun):
            if node.link:
                self._append_link(paragraph, node, report, bold=node.bold or bold, styled=styled)
            else:
                run = self._append_run(paragraph)
                self._stamp(run, self.options.body_font, styled=styled)
                _set_weight(run, node.bold or bold, node.italic)
                _set_run_text(run, node.text)
        elif isinstance(node, mp.MathRun):
            self._append_math(paragraph, node.latex, display=False, report=report, styled=styled)
        elif isinstance(node, mp.CodeSpan):
            run = self._append_run(paragraph)
            # 等宽字体是"这一格是代码"这件事本身，标题里的代码段也要保留；
            # 字号不写，跟着标题走
            _apply_font(run, self.options.code_font)
            if not styled:
                _apply_size(run, self.options.body_size)
            _set_weight(run, bold, False)
            _shade_run(run, CODE_FILL)
            _set_run_text(run, node.text)
        elif isinstance(node, mp.ImageRef):
            self._write_image(paragraph, node, report, bold=bold, styled=styled)
        else:
            report.notes.append(
                "有一种本工具还不认的行内片段（%s），已跳过" % type(node).__name__
            )

    def _write_image(
        self,
        paragraph: Paragraph,
        node: mp.ImageRef,
        report: BuildReport,
        *,
        bold: bool = False,
        styled: bool = False,
    ) -> None:
        """``![…](…)`` → 把那张图**真的嵌进** Word；拿不到就留一句带原因的话。

        尺寸只**封顶、不放大**：一张比文本区还宽的图会把版面撑破，而一张本来就只有几十像素的
        图标按原生大小摆着才是作者的意思。高也要封顶的理由是 Word 不能把一张图拆在两页上——
        比一页还高的截图不缩就直接被裁掉，剩下的那一截在纸上根本不存在。
        """
        report.images += 1
        target = mp.image_target(node.url, node.base)
        if target is None:
            reason = IMAGE_SKIP_NO_ADDRESS if not (node.url or "").strip() else IMAGE_SKIP_REMOTE
            self._skip_image(paragraph, node, report, reason, bold=bold, styled=styled)
            return
        if not mp.image_readable(target):
            self._skip_image(paragraph, node, report,
                             IMAGE_SKIP_MISSING % target, bold=bold, styled=styled)
            return
        run = self._append_run(paragraph)
        try:
            shape = run.add_picture(str(target))
        except (UnrecognizedImageError, UnexpectedEndOfFileError, InvalidImageStreamError):
            # 空 run 还挂在段里，占位那句话就写进它——换一颗 run 会把这一段的前后次序打乱
            # python-docx 认图片那三类既没有共同基类、抛出来又是空句子，所以只能一类一类点名：
            # 落到下面那一条就把英文类名拼进了中文句子。
            self._skip_image(paragraph, node, report, IMAGE_SKIP_FORMAT,
                             bold=bold, styled=styled, run=run)
            return
        except Exception as exc:
            detail = str(exc).strip()
            self._skip_image(paragraph, node, report,
                             IMAGE_SKIP_FAILED % detail if detail else IMAGE_SKIP_FORMAT,
                             bold=bold, styled=styled, run=run)
            return
        _fit_picture(shape, _content_limits(report.document))
        report.images_embedded += 1

    def _skip_image(
        self,
        paragraph: Paragraph,
        node: mp.ImageRef,
        report: BuildReport,
        reason: str,
        *,
        bold: bool = False,
        styled: bool = False,
        run: Optional[Run] = None,
    ) -> None:
        """嵌不成那一支的**唯一出口**：正文上留话、回执里记一条原因，两处说的是同一件事。"""
        report.images_skipped += 1
        run = run if run is not None else self._append_run(paragraph)
        self._stamp(run, self.options.body_font, styled=styled)
        _set_weight(run, bold, False)
        _set_run_text(run, IMAGE_PLACEHOLDER.format(name=_image_name(node), reason=reason))
        report.notes.append("图片没嵌进 Word：%s——%s" % (_image_name(node), reason))

    def _stamp(self, run: Run, font: str, *, styled: bool) -> None:
        """往 run 上写正文字体/字号；``styled=True`` 时什么都不写（段落样式那一层已经管了）。"""
        if styled:
            return
        _apply_font(run, font)
        _apply_size(run, self.options.body_size)

    # ---------------------------------------------------------------- 落笔
    def _append_run(self, paragraph: Paragraph) -> Run:
        """新建一个 ``w:r`` 接到 ``w:p`` 末尾并包成 Run。

        不用 ``paragraph.add_run()`` 是因为这一层还往同一段里接 ``w:hyperlink`` 和 OMML；
        大家一起 append，次序才等于块树里的次序。
        """
        element = OxmlElement("w:r")
        paragraph._p.append(element)
        return Run(element, paragraph)

    def _append_link(
        self, paragraph: Paragraph, node: mp.TextRun, report: BuildReport, *, bold: bool = False, styled: bool = False
    ) -> None:
        """``[文字](地址)`` → 真的 ``w:hyperlink``（蓝、下划线、点得动）。

        关系（relationship）挂在文档部件上且 ``is_external=True``——只画蓝下划线而不建关系，
        点上去没反应，那是个假链接。
        """
        url = (node.link or node.text).strip()
        try:
            r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
        except Exception as exc:  # 地址坏到连关系都建不了：退回普通文字并说明，不丢文字
            report.notes.append("超链接「%s」建不了关系（%s），已按普通文字写" % (_clip(url, 40), exc))
            url = ""

        if not url:
            run = self._append_run(paragraph)
            self._stamp(run, self.options.body_font, styled=styled)
            _set_weight(run, bold, node.italic)
            _set_run_text(run, node.text)
            return

        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), r_id)
        paragraph._p.append(link)
        for index, part in enumerate((node.text or node.link).split("\n")):
            run = Run(OxmlElement("w:r"), paragraph)
            link.append(run._element)
            self._stamp(run, self.options.body_font, styled=styled)
            _set_weight(run, bold, node.italic)
            run.font.underline = True
            run.font.color.rgb = LINK_COLOR
            if index:
                run.add_break(WD_BREAK.LINE)
            _set_run_text(run, part)
        report.links += 1

    def _append_math(
        self, paragraph: Paragraph, latex: str, *, display: bool, report: BuildReport, styled: bool = False
    ) -> None:
        """一条公式的落点：成功挂 OMML，失败退回原文 + 黄色底纹。"""
        result = self.converter.convert(latex, display=display)
        if result.ok:
            paragraph._p.append(_omml_for_word(result.omml))
            report.math_ok += 1
            return
        if self.strict:
            raise FormulaConversionError(
                "第 %d 条公式转不出来：%s\n原文：%s" % (report.math_failed + 1, result.problem, latex)
            )
        report.add_failure(result)
        run = self._append_run(paragraph)
        self._stamp(run, self.options.body_font, styled=styled)
        run.font.highlight_color = FALLBACK_HIGHLIGHT
        _set_run_text(run, "$$%s$$" % latex if display else "$%s$" % latex)


# ==========================================================================
# 直接写 OOXML 的小函数：python-docx 没有公开 API 的那些格子
# ==========================================================================


def _insert_in_order(parent, child, sequence: Sequence[str]) -> None:
    """按 schema 声明的次序把 ``child`` 插进 ``parent``。

    OOXML 的 ``w:pPr``/``w:rPr``/``w:tblPr`` 这些都是**序列类型**，子元素次序错了
    Word 会判这个包无效（不是"排版不好看"，是开不了）。python-docx 自己的 setter
    次序是对的，本层手写的那几格必须自己管。
    名单里没有的已有子元素按"排在它之后"处理，保守地不去挪动别人。
    """
    order = [qn(name) for name in sequence]
    if child.tag not in order:
        parent.append(child)
        return
    position = order.index(child.tag)
    for existing in parent:
        index = order.index(existing.tag) if existing.tag in order else -1
        if index > position:
            existing.addprevious(child)
            return
    parent.append(child)


def _get_or_add_rpr(owner):
    """Style 与 Run 的 ``w:rPr`` 都取到；取不到交回 None 而不是抛（换过模板时不至于整份崩）。"""
    element = getattr(owner, "element", None)
    if element is None:
        element = getattr(owner, "_element", None)
    if element is None:
        return None
    try:
        return element.get_or_add_rPr()
    except (AttributeError, NotImplementedError):
        return None


def _apply_font(owner, name: str) -> None:
    """给样式或 run 设西文 + 中文四格字体，并删掉主题字体属性。"""
    rpr = _get_or_add_rpr(owner)
    if rpr is None:
        return
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        _insert_in_order(rpr, fonts, RPR_SEQ)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), name)
    # 主题属性优先于显式 val：不删就等于这一格没设（模板里的 Heading 四格全是主题属性）
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attr), None)


def _apply_size(owner, point: float) -> None:
    """字号：``w:sz`` 与 ``w:szCs`` 都给（python-docx 的 ``CT_RPr`` 没声明 ``szCs``）。"""
    rpr = _get_or_add_rpr(owner)
    if rpr is None:
        return
    half_points = "%d" % max(int(round(point * 2)), 2)  # 单位是半磅
    for tag in ("w:sz", "w:szCs"):
        element = rpr.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            _insert_in_order(rpr, element, RPR_SEQ)
        element.set(qn("w:val"), half_points)


def _make_color_black(style) -> None:
    """把样式颜色钉成黑色并删掉主题色属性——模板里的 Heading 是 ``themeColor="accent1"`` 的蓝。"""
    rpr = _get_or_add_rpr(style)
    if rpr is None:
        return
    color = rpr.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        _insert_in_order(rpr, color, RPR_SEQ)
    color.set(qn("w:val"), "000000")
    for attr in ("w:themeColor", "w:themeTint", "w:themeShade"):
        color.attrib.pop(qn(attr), None)


def _set_weight(run: Run, bold: bool, italic: bool) -> None:
    """只写"要"的那一格。

    ``run.bold = False`` 落的是 ``<w:b w:val="0"/>``——那是**直接格式**，会把 Heading 样式
    自带的粗体顶掉，标题就变成"黑体不加粗"。所以否的那一格干脆不写，让样式那一层生效。
    """
    if bold:
        run.bold = True
    if italic:
        run.italic = True


def _set_run_text(run: Run, text: str) -> None:
    """写 ``w:t``，空格一个都不丢。

    ``xml:space="preserve"`` 是必须的：不写 Word 会把行首行尾的空格吃掉，
    "记号 + 空格 + 正文"那种拼接就粘在一起了。
    """
    if not text:
        return
    element = OxmlElement("w:t")
    element.set(qn("xml:space"), "preserve")
    element.text = text
    run._element.append(element)


def _set_instr_text(run: Run, text: str) -> None:
    """域代码专用：写的是 ``w:instrText`` 而不是 ``w:t``。

    ``xml:space="preserve"`` 在这里比正文里更要紧——``TOC \\o "1-3" \\h`` 的**前后各有一个空格**，
    那是 Word 分词的边界（少了前导空格会读成 ``TOC\\o``，域直接不认）。
    """
    element = OxmlElement("w:instrText")
    element.set(qn("xml:space"), "preserve")
    element.text = text
    run._element.append(element)


def _set_indent(paragraph: Paragraph, depth: int) -> None:
    if depth > 0:
        paragraph.paragraph_format.left_indent = Pt(INDENT_PT * depth)


def _shade_paragraph(paragraph: Paragraph, fill: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        _insert_in_order(ppr, shd, PPR_SEQ)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)


def _shade_run(run: Run, fill: str) -> None:
    """行内代码的底纹。``w:shd`` 排在 ``w:u``/``w:highlight`` 之后，所以调用点都放在
    设完粗体/下划线之后。"""
    rpr = run._element.get_or_add_rPr()
    shd = rpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        _insert_in_order(rpr, shd, RPR_SEQ)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)


def _quote_style(paragraph: Paragraph, depth: int) -> None:
    """引用段落的视觉：左边一条竖线 + 缩进 + 灰字。

    幂等：嵌套引用时外层会再来一遍，已经有的边框不重复加，
    已经设过的更深缩进也不改浅（内层先被写到，外层后收尾）。
    """
    ppr = paragraph._p.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is None:
        border = OxmlElement("w:pBdr")
        _insert_in_order(ppr, border, PPR_SEQ)
    if border.find(qn("w:left")) is None:
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "12")  # 1.5 磅
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), QUOTE_BORDER_COLOR)
        border.append(left)
    current = paragraph.paragraph_format.left_indent
    want = Pt(INDENT_PT * max(depth, 1))
    if current is None or current < want:
        paragraph.paragraph_format.left_indent = want
    for run in paragraph.runs:
        try:
            colored = run.font.color.rgb is not None
        except AttributeError:
            colored = True  # 主题色那一路：已经是有色了，别抢
        if not colored:
            run.font.color.rgb = QUOTE_COLOR


def _bottom_rule(paragraph: Paragraph) -> None:
    """``---`` 分隔线 → 段落底部一条横贯的细线。"""
    ppr = paragraph._p.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is None:
        border = OxmlElement("w:pBdr")
        _insert_in_order(ppr, border, PPR_SEQ)
    if border.find(qn("w:bottom")) is None:
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), BORDER_COLOR)
        border.append(bottom)


def _set_borders(table: Table) -> None:
    """表格框线自己写，不靠 ``Table Grid`` 样式名：换一份模板那一档就可能不在，
    而框线决定这张表读不读得出来。"""
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        _insert_in_order(tbl_pr, borders, TBLPR_SEQ)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement("w:" + edge)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(BORDER_SIZE_EIGHTHS))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), BORDER_COLOR)
        borders.append(element)


def _shade_cell(cell: _Cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        _insert_in_order(tc_pr, shd, TCPR_SEQ)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)


def _repeat_header_row(row) -> None:
    """表头行设 ``w:tblHeader``：这张表跨页时 Word 把这一行再印一遍。"""
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        _insert_in_order(tr_pr, header, TRPR_SEQ)


def _alignment_of(aligns: Sequence[str], index: int) -> Optional[WD_ALIGN_PARAGRAPH]:
    if index >= len(aligns):
        return None
    return ALIGNMENTS.get(aligns[index])


def _content_limits(document: Document) -> tuple[Optional[int], Optional[int]]:
    """第一页**文本区**的可用宽高（EMU），图就按这两格封顶。

    python-docx 的 ``Section`` 压根没有 ``content_width`` 这一格（1.2.0 实测），只能自己
    拿页宽减左右边距。任何一格读不出来（模板没写、继承自上一节）这一维就交回 ``None`` =
    **不封顶**，而不是猜一个数：猜小了把图缩到看不清，猜大了照样撑破版面，
    而撑破了人在 Word 里还能拖回来。
    """
    try:
        section = document.sections[0]
    except (IndexError, AttributeError):
        return None, None

    def span(total: Optional[int], *cuts: Optional[int]) -> Optional[int]:
        if total is None or any(cut is None for cut in cuts):
            return None
        return max(0, int(total) - sum(int(cut) for cut in cuts))

    return (
        span(section.page_width, section.left_margin, section.right_margin),
        span(section.page_height, section.top_margin, section.bottom_margin),
    )


def _fit_picture(shape: InlineShape, limits: tuple[Optional[int], Optional[int]]) -> None:
    """超了文本区就**按比例缩**；没超就一个字都不动（只封顶、不放大）。"""
    width, height = shape.width, shape.height
    if not width or not height:
        return
    max_width, max_height = limits
    scale = 1.0
    if max_width and width > max_width:
        scale = max_width / width
    if max_height and height * scale > max_height:
        scale = max_height / height
    if scale < 1.0:
        shape.width = Emu(int(width * scale))
        shape.height = Emu(int(height * scale))


def _basename_of_url(url: str) -> str:
    """从图片地址里挑一个认得出的名字：末段文件名，去掉查询串与锚点。"""
    cleaned = url.split("?", 1)[0].split("#", 1)[0].strip()
    if not cleaned:
        return ""
    return cleaned.replace("\\", "/").rstrip("/").rsplit("/", 1)[-1]


def _image_name(node: "mp.ImageRef") -> str:
    """占位那一格写什么：**文件名一定在**。

    ``![替代文字](地址)`` 里替代文字是给人读的描述，地址才是"去哪找回这张图"的线索；
    只写替代文字，补图的人就得回头翻 .md 才对得上是哪个文件。两个都有就都写。
    """
    alt = node.alt.strip()
    name = _basename_of_url(node.url)
    if alt and name and alt != name:
        return "%s（%s）" % (alt, name)
    return alt or name or "（没写文件名）"


def _clip(text: str, limit: int) -> str:
    return text if len(text) <= limit else text[: limit - 3] + "..."

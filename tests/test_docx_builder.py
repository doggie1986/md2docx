# -*- coding: utf-8 -*-
"""docx_builder 的自检。

跑法::

    .venv\\Scripts\\python.exe -m unittest discover -s tests -v

这一层的坑全是**"文件写得出来、Word 里却不对"**那一类：中文字体没落到 ``w:eastAsia``、
主题属性盖住显式设置、``run.bold = False`` 把样式的粗体顶掉、标题 run 上的直接格式
把 Heading 样式的黑体大字盖成宋体正文号。这些在内存里看不出来，所以断言一律读
**落盘之后**的 ``word/document.xml`` 与 ``word/styles.xml``——人在 Word 里看到的才是那一份。

本机装了 Office 才有 MML2OMML.XSL；找不到就整类跳过而不是失败。
"""

from __future__ import annotations

import re
import shutil
import struct
import sys
import tempfile
import unittest
import zipfile
import zlib
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from lxml import etree  # noqa: E402

import docx_builder as db  # noqa: E402
import formula_converter as fc  # noqa: E402
import md_parser as mp  # noqa: E402

try:
    XSLT = fc.locate_xslt()
except fc.XsltNotFoundError:  # pragma: no cover - 取决于机器上有没有 Office
    XSLT = None

M = "{%s}" % fc.MATH_NS
W = "{%s}" % fc.WORD_NS
XML_SPACE = "{%s}space" % fc.XML_NS

#: 一条 latex2mathml 认不出的式子，专门用来走"退回原文"那一路
BAD_LATEX = r"\foo \bar"


def _read(path: Path, member: str) -> str:
    with zipfile.ZipFile(str(path)) as package:
        return package.read(member).decode("utf-8")


def _xml(element: etree._Element) -> str:
    return etree.tostring(element, encoding="unicode")


def _text_of(element: etree._Element) -> str:
    """所有 ``w:t`` 拼起来（``m:t`` 不算：公式里的内容另有断言管）。"""
    return "".join(node.text or "" for node in element.findall(".//%st" % W))


def _jc(element: etree._Element) -> str | None:
    """这一段的对齐声明（``w:pPr/w:jc@val``），没有就 None。"""
    node = element.find(".//%sjc" % W)
    return None if node is None else node.get(W + "val")


def _style_block(styles: str, style_id: str) -> str:
    match = re.search(r'<w:style [^>]*w:styleId="%s".*?</w:style>' % style_id, styles, re.S)
    assert match, "styles.xml 里没有样式 %s" % style_id
    return match.group(0)


#: 72 dpi 下一个像素多少 EMU——PNG 不带 ``pHYs`` 时 python-docx 就按 72 dpi 读，
#: 所以"原生尺寸 = 像素数 × 12700"这一条在断言里当场算得出来。
EMU_PER_PIXEL = 12700


def _png(path: Path, width: int, height: int) -> Path:
    """画一张纯色 PNG 落盘。标准库就够（这个 venv 里没 Pillow，也不该为测试装一个）。

    **刻意不写 pHYs**：那正是"截图从屏幕上拷进来"最常见的样子——Word 按 72 dpi 读，
    一张 1200 像素宽的图原生就是 16.7 英寸，比 A4 的正文宽两倍多，所以封顶这一关必须有对象可测。
    """
    def chunk(tag: bytes, data: bytes) -> bytes:
        body = tag + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)

    row = b"\x00" + b"\x20\x40\x80" * width
    raw = row * height
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(b"\x89PNG\r\n\x1a\n"
                     + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
                     + chunk(b"IDAT", zlib.compress(raw, 9))
                     + chunk(b"IEND", b""))
    return path


class _Base(unittest.TestCase):
    """把"写一份 md → 落盘 → 把两份 XML 读回来"收成一次调用。"""

    maxDiff = None

    def setUp(self) -> None:
        self.converter = fc.FormulaConverter(XSLT)
        self.tmp = Path(tempfile.mkdtemp(prefix="md2docx-test-"))
        # 临时目录归临时目录：不清就会每跑一条用例在 %TEMP% 里落一个，跑测的人不会知道
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)

    def build(self, markdown: str, *, options: db.BuilderOptions | None = None, strict: bool = False) -> db.BuildReport:
        parsed = mp.parse(markdown)
        return db.DocxBuilder(self.converter, options, strict=strict).build(parsed)

    def write(self, markdown: str, *, name: str = "case.docx", **kwargs):
        """返回 ``(回执, document.xml 的 Element, styles.xml 的原文)``。"""
        report = self.build(markdown, **kwargs)
        path = report.save(self.tmp / name)
        document = etree.fromstring(_read(path, "word/document.xml").encode("utf-8"))
        return report, document, _read(path, "word/styles.xml")

    def paragraphs(self, document: etree._Element) -> list[etree._Element]:
        return document.findall(".//%sp" % W)


@unittest.skipIf(XSLT is None, "本机找不到 MML2OMML.XSL，生成用例无法真跑")
class StyleTest(_Base):
    def test_normal_style_carries_eastasia_and_size(self) -> None:
        """中文字体取的是 ``w:eastAsia``：只写 ascii/hAnsi（``font.name`` 干的就是这一件事）
        等于设了一格中文根本不读的，屏幕上就不是宋体。"""
        _, _, styles = self.write("正文一段。\n")
        normal = _style_block(styles, "Normal")
        self.assertIn('w:eastAsia="宋体"', normal)
        self.assertIn('w:ascii="宋体"', normal)
        self.assertIn('<w:sz w:val="24"/>', normal)  # 12pt = 24 半磅
        self.assertIn('<w:szCs w:val="24"/>', normal)  # python-docx 的 CT_RPr 压根没声明 szCs

    def test_heading_style_is_black_bold_and_theme_free(self) -> None:
        """标题：黑体 + 正文号往上加档 + 加粗 + 黑色，而且**主题属性要删**。
        模板里 Heading 的 ``w:rFonts`` 四格全是 ``*Theme``、``w:color`` 带 ``themeColor``，
        规范规定主题属性优先于显式 val——只 set 不删就等于没设，标题照旧是蓝色 Calibri Light。"""
        _, _, styles = self.write("# 标题\n")
        block = _style_block(styles, "Heading1")
        self.assertIn('w:eastAsia="黑体"', block)
        self.assertIn('<w:sz w:val="34"/>', block)  # 12 + 5 = 17pt
        self.assertIn("<w:b/>", block)
        self.assertIn('w:val="000000"', block)
        self.assertNotIn("Theme", block)
        self.assertNotIn("themeColor", block)

    def test_heading_runs_do_not_carry_direct_font_or_size(self) -> None:
        """这一条是 2026-10-14 真开 Word 抓到的缺陷：段落挂了 Heading 样式，run 上又盖了
        正文字体与字号 = 直接格式盖过样式，需求里"标题用黑体"当场失效。"""
        _, document, _ = self.write("# 一级标题\n\n## 二级标题\n\n正文一段。\n")
        headings = [p for p in self.paragraphs(document) if p.find("%spPr/%spStyle" % (W, W)) is not None]
        self.assertEqual(
            [p.find("%spPr/%spStyle" % (W, W)).get(W + "val") for p in headings],
            ["Heading1", "Heading2"],
        )
        for paragraph in headings:
            for run in paragraph.findall("%sr" % W):
                self.assertIsNone(run.find("%srPr" % W), "标题的 run 上不该有任何直接格式")
        body = self.paragraphs(document)[-1]
        self.assertIsNotNone(body.find("%sr/%srPr/%srFonts" % (W, W, W)), "正文 run 仍自带字体声明")

    def test_inline_bold_inside_heading_is_the_only_direct_format(self) -> None:
        """标题里 ``**加粗**`` 仍要认：那是内容的一部分，不是把整段的样式盖掉。

        "一级标题里有 **加粗** 二字" 切成三段文字，所以是三个 run；带直接格式的只有中间那个，
        而且它只带 ``<w:b/>``——字体字号仍归 Heading 样式管。
        """
        _, document, _ = self.write("# 一级标题里有 **加粗** 二字\n")
        runs = self.paragraphs(document)[0].findall("%sr" % W)
        self.assertEqual([_text_of(run) for run in runs], ["一级标题里有 ", "加粗", " 二字"])
        with_format = [run for run in runs if run.find("%srPr" % W) is not None]
        self.assertEqual(len(with_format), 1, "标题里只有加粗那一处该带直接格式")
        self.assertIsNotNone(with_format[0].find("%srPr/%sb" % (W, W)))
        self.assertIsNone(with_format[0].find("%srPr/%srFonts" % (W, W)), "只有粗体，字体字号仍归样式")

    def test_options_flow_into_styles(self) -> None:
        """命令行的 --body-font / --body-size / --heading-font 走的就是 BuilderOptions，
        它得真的改到 styles.xml；10.5pt 落成 21 半磅（不许有 float 误差）。"""
        options = db.BuilderOptions(body_font="微软雅黑", body_size=10.5, heading_font="楷体")
        _, _, styles = self.write("正文\n\n# 标题\n", options=options)
        self.assertIn('w:eastAsia="微软雅黑"', styles)
        self.assertIn('<w:sz w:val="21"/>', _style_block(styles, "Normal"))
        heading = _style_block(styles, "Heading1")
        self.assertIn('w:eastAsia="楷体"', heading)
        self.assertIn('<w:sz w:val="31"/>', heading)  # 10.5 + 5 = 15.5pt = 31 半磅

    def test_heading_size_is_relative_to_body_size(self) -> None:
        """标题字号是"正文号 + 增量"：写死绝对号的话 --body-size 一改标题就乱了。"""
        self.assertEqual(db.BuilderOptions(body_size=12).heading_size(1), 17.0)
        self.assertEqual(db.BuilderOptions(body_size=10.5).heading_size(2), 13.5)
        self.assertEqual(db.BuilderOptions(body_size=12).heading_size(6), 12.0)

    def test_missing_heading_style_falls_back_to_bold_text(self) -> None:
        """换过模板、样式表里压根没有 Heading 3 时不该崩，也不该把标题写成看不出是标题的一段。"""
        with mock.patch.object(db.DocxBuilder, "_style", staticmethod(lambda document, name: None)):
            report, document, _ = self.write("### 三级标题\n")
        self.assertTrue(any("Heading 3" in note for note in report.notes), report.notes)
        run = self.paragraphs(document)[0].find("%sr" % W)
        self.assertIsNotNone(run.find("%srPr/%sb" % (W, W)), "样式没挂上时用加粗把标题撑出来")
        self.assertIsNotNone(run.find("%srPr/%srFonts" % (W, W)), "样式没了，字体就只能写在 run 上")


@unittest.skipIf(XSLT is None, "本机找不到 MML2OMML.XSL，生成用例无法真跑")
class MathPlacementTest(_Base):
    def test_inline_math_lands_in_paragraph_as_omath(self) -> None:
        report, document, _ = self.write("灵敏度 $S = \\beta I_0$ 一句话。\n")
        self.assertEqual((report.math_ok, report.math_failed), (1, 0))
        self.assertEqual(len(document.findall(".//%soMath" % M)), 1)
        self.assertEqual(len(document.findall(".//%soMathPara" % M)), 0, "行内公式不该外包 oMathPara")

    def test_display_math_wraps_and_centers(self) -> None:
        """块级公式外包 ``m:oMathPara``——Word 认这一档才会把 ``\\sum`` 的上下限摆到正上/正下方。

        写成独占一行的 ``$$…$$``（解析层的口径：混在文字里的那一对 ``$$`` 算行内，不外包
        ``oMathPara``，也不居中）。
        """
        _, document, _ = self.write("$$P = UI\\cos\\varphi$$\n")
        para = document.find(".//%soMathPara" % M)
        self.assertIsNotNone(para)
        self.assertEqual(_jc(para.getparent()), "center")

    def test_text_formula_link_formula_order_is_preserved(self) -> None:
        """一行里"文字—公式—文字—链接—文字—公式—文字"的次序由块树决定。
        OMML 与 w:hyperlink 都是手动 append 的；走 python-docx 的 setter 会按各自的插入规则重排。"""
        _, document, _ = self.write("见 $a$ 与 [手册](https://example.com/a) 及 $b$。\n")
        tags = [
            child.tag.replace(M, "m:").replace(W, "w:")
            for child in document.find(".//%sp" % W)
            if child.tag != W + "pPr"
        ]
        self.assertEqual(tags, ["w:r", "m:oMath", "w:r", "w:hyperlink", "w:r", "m:oMath", "w:r"])

    def test_math_runs_declare_cambria_math_and_preserve_spaces(self) -> None:
        """公式的字形声明挂在**每个** ``m:r`` 上；``m:t`` 少了 ``xml:space`` 空格会被吃掉。"""
        _, document, _ = self.write("行内 $a + b$。\n")
        runs = document.findall(".//%soMath/%sr" % (M, M))
        self.assertTrue(runs)
        for run in runs:
            fonts = run.find("%srPr/%srFonts" % (W, W))
            self.assertIsNotNone(fonts, "公式的 m:r 没带 w:rFonts，Word 会换成默认字体")
            self.assertEqual(fonts.get(W + "ascii"), "Cambria Math")
        for node in document.findall(".//%st" % M):
            self.assertEqual(node.get(XML_SPACE), "preserve")

    def test_formula_in_a_table_cell_stays_inside_that_cell(self) -> None:
        """需求点名的一条：单元格里的公式要转对，而且不能跑到表外面去。"""
        report, document, _ = self.write("| 名称 | 表达式 |\n| --- | --- |\n| 功率 | $P = UI$ |\n")
        self.assertEqual(report.math_ok, 1)
        math = document.find(".//%soMath" % M)
        self.assertIsNotNone(math)
        ancestors = set()
        node = math
        while node is not None:
            ancestors.add(node.tag)
            node = node.getparent()
        self.assertIn(W + "tc", ancestors, "公式的祖先里没有单元格——它跑到表外面去了")

    def test_display_math_inside_cell_does_not_escape_the_table(self) -> None:
        """表格里写 ``\\[…\\]``：独占一段并居中，但段落**仍在单元格里**——
        跑出去等于表格少了一格，Word 里看到的是一张缺列的表。"""
        _, document, _ = self.write("| 名称 | 表达式 |\n| --- | --- |\n| 线损率 | \\[ \\eta = \\dfrac{\\Delta A}{A} \\] |\n")
        para = document.find(".//%soMathPara" % M)
        self.assertIsNotNone(para)
        self.assertEqual(para.getparent().getparent().tag, W + "tc")
        self.assertEqual(_jc(para.getparent()), "center")

    def test_delimiters_never_survive_into_the_document(self) -> None:
        """四种定界符都不该在成品正文里出现（故意写坏、退回原文那一路除外）。"""
        _, document, _ = self.write(r"行内 $a+b$，块级" + "\n\n$$c+d$$\n\n" + r"\(e+f\) 与 \[g+h\]" + "\n")
        text = _text_of(document)
        self.assertNotIn("$", text)
        self.assertNotIn("\\(", text)
        self.assertNotIn("\\[", text)


@unittest.skipIf(XSLT is None, "本机找不到 MML2OMML.XSL，生成用例无法真跑")
class FailureTest(_Base):
    def test_bad_formula_keeps_source_text_and_highlight(self) -> None:
        """默认档：坏公式退回原文 + 黄色底纹，整份文档照出——
        几百条公式的稿子不该为一条崩掉（业务方 2026-10-14 定的口径）。"""
        report, document, _ = self.write("这里 $%s$ 写坏了。\n" % BAD_LATEX)
        self.assertEqual((report.math_ok, report.math_failed), (0, 1))
        self.assertEqual(len(report.failures), 1)
        self.assertIn(BAD_LATEX, report.failures[0].failure_report())
        self.assertIn(BAD_LATEX, _text_of(document))
        self.assertIn('<w:highlight w:val="yellow"/>', _xml(document))

    def test_strict_raises_and_says_which_formula(self) -> None:
        with self.assertRaises(fc.FormulaConversionError) as caught:
            self.build("这里 $%s$ 写坏了。\n" % BAD_LATEX, strict=True)
        self.assertIn(BAD_LATEX, str(caught.exception))

    def test_good_formulas_still_convert_next_to_a_bad_one(self) -> None:
        report, document, _ = self.write("对的 $a+b$，坏的 $%s$，再一对 $c+d$。\n" % BAD_LATEX)
        self.assertEqual((report.math_ok, report.math_failed), (2, 1))
        self.assertEqual(len(document.findall(".//%soMath" % M)), 2)

    def test_receipt_lists_failures_and_counts(self) -> None:
        report = self.build("坏的 $%s$，好的 $x$。\n" % BAD_LATEX)
        lines = "\n".join(report.lines())
        self.assertIn("公式 2 条（转成 1，失败 1）", lines)
        self.assertIn(BAD_LATEX, lines)
        self.assertEqual(report.math_total, 2)


@unittest.skipIf(XSLT is None, "本机找不到 MML2OMML.XSL，生成用例无法真跑")
class TableTest(_Base):
    def test_header_row_is_shaded_bold_and_repeats(self) -> None:
        _, document, _ = self.write("| 名称 | 数量 |\n|---|---|\n| 导线 | 3 |\n")
        first = document.findall(".//%str" % W)[0]
        self.assertIsNotNone(first.find("%strPr/%stblHeader" % (W, W)), "跨页不重复表头")
        self.assertIn('w:fill="D9D9D9"', _xml(first))
        self.assertIsNotNone(first.find(".//%sr/%srPr/%sb" % (W, W, W)))

    def test_column_alignment_reaches_every_cell(self) -> None:
        _, document, _ = self.write("| 左 | 中 | 右 |\n|:---|:---:|---:|\n| 1 | 2 | 3 |\n")
        for row in document.findall(".//%str" % W):
            self.assertEqual([_jc(cell) for cell in row.findall("%stc" % W)], ["left", "center", "right"])

    def test_ragged_row_is_padded_and_warned(self) -> None:
        """某行少一格：补齐成表头的列数（否则 Word 里那张表的形状就变了），并念一句提醒。"""
        report, document, _ = self.write("| a | b | c |\n|---|---|---|\n| 1 | 2 |\n")
        cells = document.findall(".//%str" % W)[-1].findall("%stc" % W)
        self.assertEqual(len(cells), 3)
        self.assertEqual([_text_of(c) for c in cells], ["1", "2", ""])
        self.assertTrue(any("补齐" in w for w in report.parse_warnings), report.parse_warnings)

    def test_two_adjacent_tables_are_not_merged(self) -> None:
        """两张表中间必须有段落，否则 Word 把它们并成一张（列数对得上就并得更干脆），
        并完没人看得出这里原本有两张表。"""
        _, document, _ = self.write("| a |\n|---|\n| 1 |\n\n| b |\n|---|\n| 2 |\n")
        body = [child.tag for child in document.find("%sbody" % W)]
        tables = [i for i, tag in enumerate(body) if tag == W + "tbl"]
        self.assertEqual(len(tables), 2)
        self.assertTrue(any(body[i] == W + "p" for i in range(tables[0] + 1, tables[1])))

    def test_document_does_not_end_on_a_table(self) -> None:
        _, document, _ = self.write("| a |\n|---|\n| 1 |\n")
        body = [child.tag for child in document.find("%sbody" % W) if child.tag != W + "sectPr"]
        self.assertEqual(body[-1], W + "p")

    def test_bold_never_writes_an_explicit_off_switch(self) -> None:
        """``run.bold = False`` 落的是 ``<w:b w:val="0"/>``——那是**直接格式**，
        会把 Heading 样式自带的粗体顶掉。全文只许出现 ``<w:b/>`` 这一种写法。"""
        _, document, _ = self.write("# 标题里有 **加粗** 和小体 *斜体*\n\n普通一段。\n")
        xml = _xml(document)
        self.assertNotIn('<w:b w:val="0"', xml)
        self.assertNotIn('<w:i w:val="0"', xml)
        self.assertIn("<w:b/>", xml)
        self.assertIn("<w:i/>", xml)


@unittest.skipIf(XSLT is None, "本机找不到 MML2OMML.XSL，生成用例无法真跑")
class PackageTest(_Base):
    def test_package_reopens_and_has_no_ns0_prefix(self) -> None:
        """schema 次序写错的包 python-docx 写得出、Word 打不开；重开一次是最省事的把关。
        ``ns0:`` 是 lxml 的自动前缀：Word 认 URI 不认前缀，但对账的人要解压看 document.xml，
        ``<ns0:oMathPara>`` 与 ``<m:oMathPara>`` 差一口气。"""
        report = self.build("| a | b |\n|---|---|\n| $x^2$ | 文本 |\n\n$$y \\ne x$$\n")
        path = report.save(self.tmp / "reopen.docx")
        from docx import Document  # 读开一次：包坏了这里就抛

        self.assertTrue(Document(str(path)).paragraphs)
        document = _read(path, "word/document.xml")
        self.assertNotIn("ns0:", document)
        self.assertIn("<m:oMathPara>", document)
        self.assertIn("<w:rFonts", document)

    def test_count_omml_matches_the_report(self) -> None:
        """回执说转成几条，文件里就得有几个 m:oMath——这才是"原生对象而不是图片"的证据。

        ``count_omml`` 数的是**全部** ``m:oMath``，显示式外包的 ``m:oMathPara`` 里那一个也算，
        所以三个数是这样对的：2 条行内 + 1 条显示式 = 3 个 m:oMath，其中 1 个在 m:oMathPara 里。
        """
        report = self.build("行内 $a$ 与 $b$，块级\n\n$$c$$\n")
        path = report.save(self.tmp / "count.docx")
        self.assertEqual(db.count_omml(path), (3, 1))
        self.assertEqual(report.math_ok, 3)

    def test_link_is_a_real_hyperlink_with_external_relationship(self) -> None:
        """只画蓝下划线而不建关系是个假链接，点上去没反应。"""
        report, document, _ = self.write("[口径说明](https://example.com/spec)\n")
        self.assertEqual(report.links, 1)
        self.assertIn("<w:hyperlink", _xml(document))
        with zipfile.ZipFile(str(self.tmp / "case.docx")) as package:
            rels = package.read("word/_rels/document.xml.rels").decode("utf-8")
        self.assertIn('Target="https://example.com/spec"', rels)
        self.assertIn('TargetMode="External"', rels)

    def test_save_creates_missing_parent_directories(self) -> None:
        """命令行指到没建过的子目录是很常见的写法。"""
        report = self.build("一段。\n")
        path = report.save(self.tmp / "nested" / "deeper" / "out.docx")
        self.assertTrue(path.is_file())


@unittest.skipIf(XSLT is None, "本机找不到 MML2OMML.XSL，生成用例无法真跑")
class OtherBlocksTest(_Base):
    def test_code_block_is_monospace_shaded_and_hides_dollar_signs(self) -> None:
        """代码块里的 ``$`` 在解析层就挡在公式识别之外；整块等宽字体 + 浅灰底。"""
        report, document, _ = self.write("```python\nrate = price * 2  # $/kg\n```\n")
        self.assertEqual((report.math_ok, report.math_failed), (0, 0))
        xml = _xml(document)
        self.assertIn("Consolas", xml)
        self.assertIn('w:fill="F2F2F2"', xml)
        self.assertIn("$/kg", _text_of(document))

    def test_inline_code_keeps_its_font(self) -> None:
        _, document, _ = self.write("命令 `md2docx.py a.md b.docx` 一行。\n")
        self.assertIn("Consolas", _xml(document))

    def test_list_markers_are_literal_text(self) -> None:
        """自动编号要靠 ``word/numbering.xml`` 里的编号定义，默认模板那套不完整
        （实测开出来是几行光秃秃的段落）。写死记号换来"屏幕上的号就是稿子里的那几个数"。"""
        _, document, _ = self.write("- 甲\n- 乙\n\n1. 第一\n2. 第二\n")
        texts = [_text_of(p) for p in self.paragraphs(document)]
        self.assertEqual(texts[:2], ["• 甲", "• 乙"])
        self.assertEqual(texts[2:4], ["1. 第一", "2. 第二"])

    def test_ordered_list_honours_start_and_nesting(self) -> None:
        """从 3 起编号就按 3 写；子项另起一段并多缩进一层。"""
        _, document, _ = self.write("3. 丙\n   - 子甲\n")
        paragraphs = self.paragraphs(document)
        self.assertEqual([_text_of(p) for p in paragraphs], ["3. 丙", "• 子甲"])
        indents = [p.find("%spPr/%sind" % (W, W)) for p in paragraphs]
        self.assertTrue(all(node is not None for node in indents), indents)
        lefts = [int(node.get(W + "left")) for node in indents]
        self.assertGreater(lefts[1], lefts[0], "子项没有比父项多缩进一层")

    def test_quote_block_is_indented_with_a_left_rule(self) -> None:
        """不用 Word 的「Intense Quote」样式——那一档在换过模板的机器上不保证存在；
        缩进和边框是自己写的，走到哪台机器都长一个样。"""
        _, document, _ = self.write("> 引用一句\n")
        xml = _xml(document)
        self.assertIn("<w:pBdr>", xml)
        self.assertIn("<w:left ", xml)
        self.assertIn("引用一句", _text_of(document))

    def test_thematic_break_writes_a_bottom_rule(self) -> None:
        _, document, _ = self.write("上一段\n\n---\n\n下一段\n")
        self.assertIn("<w:bottom ", _xml(document))

    def test_empty_input_still_produces_a_document_and_a_note(self) -> None:
        """输入是空的：交回一份有段落的东西 + 一句说明，而不是抛错或一份零段文档。"""
        report, document, _ = self.write("\n\n")
        self.assertEqual(report.blocks, 0)
        self.assertTrue(report.notes)
        self.assertTrue(self.paragraphs(document))


@unittest.skipIf(XSLT is None, "本机找不到 MML2OMML.XSL，生成用例无法真跑")
class ImageEmbedTest(_Base):
    """``![…](…)`` 指向的图**真的嵌进** Word：读盘、按文本区封顶、嵌不成的每一处留一句原因。

    这一族照本文件开头的口径：断言一律读**落盘之后**的那一份包——``word/media/`` 里有没有那个
    文件、``wp:extent`` 写的是多少，才是人在 Word 里看到的东西；内存里的对象不算证据。
    """

    def build_md(self, markdown: str, name: str = "a.md") -> db.BuildReport:
        """写成一份 .md 再走 :func:`md_parser.parse_file`。

        必须走这一条：图片"去哪找"那一格（``ImageRef.base``）是 ``parse_file`` 盖上去的，
        命令行与界面走的都是它；直接 ``mp.parse`` 的结果 base 是 None，会按当前工作目录找——
        那是另一条路（校验那边才有"调用方给的目录"这一档）。
        """
        path = self.tmp / name
        path.write_text(markdown, encoding="utf-8")
        return db.DocxBuilder(self.converter).build(mp.parse_file(path))

    def reopen(self, report: db.BuildReport, name: str = "a.docx"):
        """落盘后重开一次，交回 ``(内嵌图, 媒体文件名列表, document.xml 的全部文本)``。"""
        from docx import Document  # 读开一次：包坏了这里就抛

        path = report.save(self.tmp / name)
        with zipfile.ZipFile(str(path)) as package:
            media = sorted(n for n in package.namelist() if n.startswith("word/media/"))
        document = Document(str(path))
        text = _text_of(etree.fromstring(_read(path, "word/document.xml").encode("utf-8")))
        return list(document.inline_shapes), media, text, document

    def test_a_local_png_becomes_a_media_part(self) -> None:
        """这是这一轮要的那件事本身：文件在，图就得在包里，而不是留一句占位。"""
        _png(self.tmp / "tower.png", 60, 40)
        report = self.build_md("![杆塔示意](tower.png)\n")
        shapes, media, text, _doc = self.reopen(report)
        self.assertEqual(["word/media/image1.png"], media)
        self.assertEqual(1, len(shapes))
        self.assertNotIn("图片没嵌进去", text)
        self.assertEqual((1, 1, 0),
                         (report.images, report.images_embedded, report.images_skipped))
        self.assertEqual([], [n for n in report.notes if "图片" in n])

    def test_a_picture_that_fits_keeps_its_native_size(self) -> None:
        """只封顶、**不放大**：一张图标按原生大小摆着才是作者的意思，放大就是糊。"""
        _png(self.tmp / "fit.png", 60, 40)
        report = self.build_md("![小图](fit.png)\n", "fit.md")
        shapes, _media, _text, _doc = self.reopen(report, "fit.docx")
        self.assertEqual(60 * EMU_PER_PIXEL, shapes[0].width)
        self.assertEqual(40 * EMU_PER_PIXEL, shapes[0].height)

    def test_a_picture_wider_than_the_text_area_is_capped(self) -> None:
        """超了文本区就按比例缩：多出来的那一截在纸上根本不存在，而裁切不会告诉你。"""
        _png(self.tmp / "wide.png", 1200, 300)
        report = self.build_md("![宽截图](wide.png)\n", "wide.md")
        shapes, _media, _text, document = self.reopen(report, "wide.docx")
        shape = shapes[0]
        limit_width, limit_height = db._content_limits(document)
        self.assertEqual(limit_width, shape.width, "封顶要正好卡在文本区宽")
        self.assertLessEqual(shape.height, limit_height)
        self.assertAlmostEqual(4.0, shape.width / shape.height, places=3)
        # extent 是**两处**都要写对（wp:extent 管占位、a:ext 管拉伸后的实际大小），
        # 只读回一个数看不出没同步——这里直接读文件原文。
        xml = _read(self.tmp / "wide.docx", "word/document.xml")
        self.assertIn('cx="%d"' % limit_width, xml)

    def test_a_picture_taller_than_the_page_is_capped_by_height(self) -> None:
        """Word 不能把一张图拆在两页上，所以高也要封顶——只卡宽的话长截图照样被裁掉一截。"""
        _png(self.tmp / "tall.png", 100, 5000)
        report = self.build_md("![长截图](tall.png)\n", "tall.md")
        shapes, _media, _text, document = self.reopen(report, "tall.docx")
        shape = shapes[0]
        limit_width, limit_height = db._content_limits(document)
        self.assertLessEqual(shape.height, limit_height)
        self.assertLess(shape.width, limit_width, "按高度缩完之后宽也在区内")
        self.assertAlmostEqual(100 / 5000, shape.width / shape.height, places=3)

    def test_a_missing_file_leaves_the_reason_and_the_place_it_looked(self) -> None:
        """报"找不到"必须把**查过的位置**念出来：只说没找到，人就得自己算相对路径落在哪。"""
        report = self.build_md("![杆塔示意](images/nope.png)\n", "miss.md")
        shapes, media, text, _doc = self.reopen(report, "miss.docx")
        self.assertEqual([], media)
        self.assertEqual([], shapes)
        self.assertIn("这台机器上没找到这个文件", text)
        # 比 resolve 过的那一份：``self.tmp`` 是 tempfile 给的 8.3 短名，
        # 而消息里念的是长路径——按原样比就成了一局"路径没错、断言红了"。
        self.assertIn(str((self.tmp / "images" / "nope.png").resolve()), text)
        self.assertIn("杆塔示意", text)
        self.assertEqual(1, report.images_skipped)
        self.assertEqual(1, len([n for n in report.notes if "图片没嵌进 Word" in n]))

    def test_a_remote_address_says_this_tool_does_not_fetch_it(self) -> None:
        """远程地址**不能**念成"找不到这个文件"——那是本机的写法，人听了会去磁盘上翻。"""
        report = self.build_md("![远程图](https://example.com/a.png)\n", "remote.md")
        _shapes, media, text, _doc = self.reopen(report, "remote.docx")
        self.assertEqual([], media)
        self.assertIn("这不是本机的文件", text)
        self.assertNotIn("没找到这个文件", text)

    def test_a_file_that_is_not_a_picture_is_not_called_missing(self) -> None:
        """文件在、但不是图：句子要说"读不动 / Word 能放哪几档"，而不是"没找到"。"""
        (self.tmp / "note.txt").write_text("我不是图片", encoding="utf-8")
        report = self.build_md("![说明文字](note.txt)\n", "notimage.md")
        _shapes, media, text, _doc = self.reopen(report, "notimage.docx")
        self.assertEqual([], media)
        self.assertIn("这个文件读不动", text)
        self.assertIn("PNG", text)
        self.assertNotIn("没找到这个文件", text)

    def test_an_empty_file_is_not_reported_with_an_english_class_name(self) -> None:
        """零字节文件那一类 python-docx 抛的是另一个异常、且**不带句子**：硬拼就成了
        ``这个文件读不动（InvalidImageFileError）``——同一张稿子里冒英文类名是这一路被
        验收退回过的那件事。"""
        (self.tmp / "empty.png").write_bytes(b"")
        report = self.build_md("![空文件](empty.png)\n", "empty.md")
        _shapes, media, text, _doc = self.reopen(report, "empty.docx")
        self.assertEqual([], media)
        self.assertIn("这个文件读不动", text)
        self.assertNotIn("Error", text)

    def test_a_truncated_picture_is_not_reported_with_an_english_class_name(self) -> None:
        """截断的 PNG 抛的是第三类异常（``UnexpectedEndOfFileError``，同样不带句子）：
        只认 ``UnrecognizedImageError`` 的话这一条会落到兜底那一句、把类名念出来。"""
        _png(self.tmp / "half.png", 24, 12)
        path = self.tmp / "half.png"
        path.write_bytes(path.read_bytes()[: 12])
        report = self.build_md("![只有一半](half.png)\n", "half.md")
        _shapes, media, text, _doc = self.reopen(report, "half.docx")
        self.assertEqual([], media)
        self.assertIn("这个文件读不动", text)
        self.assertNotIn("Error", text)

    def test_the_receipt_gives_both_counts(self) -> None:
        """嵌进几条、没成几条**分开报**：只念"图片 3 处"会把"有一张没找到"演成没事。"""
        _png(self.tmp / "ok.png", 20, 10)
        report = self.build_md("![好](ok.png)\n\n![坏](https://x/y.png)\n\n"
                               "![没](gone.png)\n", "three.md")
        self.assertEqual((3, 1, 2),
                         (report.images, report.images_embedded, report.images_skipped))
        self.assertIn("图片 3 处（嵌进 1，没成 2）", report.lines()[0])

    def test_two_source_directories_each_find_their_own_picture(self) -> None:
        """合并稿的常态：两个文件夹里各有一张**同名**的图，各自回各自那一格找。

        base 住在节点上而不是构造参数上就是为了这一条——一个目录管全场的话，
        第二个文件的图会被判成"找不到"，而那份文件里明明写着它。
        两张画成**不同大小**：字节完全相同的两份图 Word 会并成一个媒体文件（按校验和认），
        那一格就在测"两个目录各读到了自己的那一份"以外的事了。
        """
        for index, folder in enumerate(("第一册", "第二册")):
            _png(self.tmp / folder / "same.png", 30 + 10 * index, 20)
            (self.tmp / folder / "body.md").write_text("![同名的图](same.png)\n",
                                                       encoding="utf-8")
        first = mp.parse_file(self.tmp / "第一册" / "body.md")
        second = mp.parse_file(self.tmp / "第二册" / "body.md")
        merged = mp.ParseResult(blocks=first.blocks + second.blocks)
        report = db.DocxBuilder(self.converter).build(merged)
        shapes, media, text, _doc = self.reopen(report, "merged.docx")
        self.assertEqual(2, len(shapes))
        self.assertEqual(2, len(media))
        self.assertNotIn("图片没嵌进去", text)
        self.assertEqual(2, report.images_embedded)

    def test_a_picture_inside_a_table_cell_is_embedded_too(self) -> None:
        """表格格子里的图片也走同一条：那一格是"这一行有图"的常见写法，不能只认正文。"""
        _png(self.tmp / "cell.png", 40, 20)
        report = self.build_md("| 名称 | 图 |\n|---|---|\n| 截面 | ![截面](cell.png) |\n",
                               "cell.md")
        shapes, media, _text, _doc = self.reopen(report, "cell.docx")
        self.assertEqual(1, len(media))
        self.assertEqual(1, len(shapes))
        self.assertEqual(1, report.images_embedded)


if __name__ == "__main__":  # pragma: no cover
    unittest.main(verbosity=2)
